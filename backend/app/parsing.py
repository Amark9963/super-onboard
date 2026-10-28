"""Deterministic, bounded parsing. No OCR, network fetches, or AI calls."""

import io
import logging
import multiprocessing
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath

from app.config import settings

TYPES = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".html": "text/html",
    ".htm": "text/html",
}
MAX_TEXT = 2_000_000
MAX_BLOCKS = 5000


class ParseIssue(ValueError):
    pass


_BLOCK_TAGS = {
    "p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "td", "th", "blockquote", "dt", "dd",
    "pre", "figcaption", "caption", "summary", "address",
}  # fmt: skip
_SKIP_TAGS = {"script", "style", "template", "noscript", "svg", "iframe", "object", "embed"}
_BREAK_TAGS = {
    "div", "section", "article", "header", "footer", "nav", "main", "aside", "ul", "ol", "dl",
    "form", "fieldset", "thead", "tbody", "tfoot", "figure", "details", "menu",
}  # fmt: skip
_VOID_TAGS = {"br", "hr", "img", "input", "meta", "link", "base", "col", "area", "wbr", "source"}


def _hidden(attributes: dict) -> bool:
    """Text a human reader cannot see is not evidence (and is a prompt-injection channel)."""
    style = attributes.get("style", "").replace(" ", "").lower()
    return (
        "hidden" in attributes
        or attributes.get("aria-hidden", "").lower() == "true"
        or "display:none" in style
        or "visibility:hidden" in style
    )


class _HtmlBlocks(HTMLParser):
    """Visible text as ordered blocks with heading paths. Nothing in the page is executed."""

    def __init__(self, add):
        super().__init__(convert_charrefs=True)
        self.add, self.title = add, ""
        self.headings: list[str] = []
        self.skip = 0  # depth inside skipped/hidden subtrees
        self.hidden_stack: list[str] = []
        self.in_title = False
        self.buffer: list[str] = []
        self.current: str | None = None
        self.tables = 0
        self.row = self.column = 0
        self.number = 0

    def _flush(self):
        text = " ".join("".join(self.buffer).split())
        self.buffer = []
        tag, self.current = self.current, None
        if not text:
            return
        self.number += 1
        locator = {"section_path": list(self.headings), "element": tag or "text",
                   "block_number": self.number}  # fmt: skip
        if tag in {"td", "th"}:
            locator.update(table=self.tables, row=self.row, column=self.column)
        if tag and tag[0] == "h" and len(tag) == 2 and tag[1].isdigit():
            level = int(tag[1])
            self.headings = self.headings[: max(level - 1, 0)] + [text[:200]]
            locator["section_path"] = list(self.headings)
        self.add(text, locator)

    def handle_starttag(self, tag, attrs):
        attributes = {k: (v or "") for k, v in attrs}
        if self.skip:
            if tag not in _VOID_TAGS:
                self.hidden_stack.append(tag)
                self.skip += 1
            return
        if tag == "title":
            self.in_title = True
            return
        if tag in _SKIP_TAGS or _hidden(attributes):
            if tag not in _VOID_TAGS:
                self.hidden_stack.append(tag)
                self.skip += 1
            return
        if tag == "table":
            self._flush()
            self.tables, self.row, self.column = self.tables + 1, 0, 0
        elif tag == "tr":
            self._flush()
            self.row, self.column = self.row + 1, 0
        if tag in {"td", "th"}:
            self.column += 1
        if tag in _BLOCK_TAGS:
            self._flush()
            self.current = tag
        elif tag in _BREAK_TAGS:
            self._flush()
        elif tag in {"br"}:
            self.buffer.append(" ")

    def handle_endtag(self, tag):
        if self.skip:
            if self.hidden_stack and self.hidden_stack[-1] == tag:
                self.hidden_stack.pop()
                self.skip -= 1
            return
        if tag == "title":
            self.in_title = False
        elif tag in _BLOCK_TAGS or tag in _BREAK_TAGS:
            self._flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.in_title:
            self.title += data
        else:
            self.buffer.append(data)

    def close(self):
        super().close()
        self._flush()


def validate_file(filename: str, content: bytes) -> str:
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in TYPES:
        raise ParseIssue("unsupported_file_type")
    if not content:
        raise ParseIssue("empty_file")
    if len(content) > settings.max_upload_bytes:
        raise ParseIssue("file_too_large")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise ParseIssue("file_signature_mismatch")
    if extension in {".docx", ".xlsx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 40_000_000:
                    raise ParseIssue("archive_expansion_limit")
                if any(
                    e.file_size > 20_000_000 or e.file_size / max(e.compress_size, 1) > 200
                    for e in entries
                ):
                    raise ParseIssue("archive_expansion_limit")
                expected = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
                if (
                    expected not in archive.namelist()
                    or "[Content_Types].xml" not in archive.namelist()
                ):
                    raise ParseIssue("file_signature_mismatch")
                if any("vbaproject" in e.filename.lower() for e in entries):
                    raise ParseIssue("macros_not_supported")
        except zipfile.BadZipFile:
            raise ParseIssue("invalid_office_archive") from None
    if extension in {".html", ".htm"} and b"\x00" in content:
        raise ParseIssue("binary_text_not_supported")
    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ParseIssue("text_requires_utf8") from None
        if any(ord(c) < 32 and c not in "\n\r\t" for c in text):
            raise ParseIssue("binary_text_not_supported")
    return TYPES[extension]


def parse_document(filename: str, data: bytes, transcript=False) -> dict:
    validate_file(filename, data)
    extension = PurePosixPath(filename).suffix.lower()
    blocks = []
    total = 0
    issue = None
    title = None

    def add(text, locator):
        nonlocal total
        text = str(text).strip()
        if not text:
            return
        total += len(text)
        if total > MAX_TEXT or len(blocks) >= MAX_BLOCKS:
            raise ParseIssue("extracted_content_limit")
        blocks.append({"text": text, "locator": locator, "position": len(blocks)})

    if extension == ".txt":
        for line_number, line in enumerate(data.decode("utf-8-sig").splitlines(), 1):
            locator = {"line_start": line_number, "line_end": line_number}
            if transcript:
                match = re.match(r"^\s*\[?(\d{1,2}:\d{2}(?::\d{2})?)\]?\s", line)
                if match:
                    values = [int(part) for part in match[1].split(":")]
                    if all(value < 60 for value in values[1:]):
                        seconds = sum(value * 60**i for i, value in enumerate(reversed(values)))
                        locator["timestamp_start"] = seconds
                        locator["timestamp_label"] = match[1]
            add(line, locator)
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseIssue("encrypted_pdf_not_supported")
        if len(reader.pages) > 250:
            raise ParseIssue("page_limit")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                issue = "pages_without_text_ocr_required"
            add(text, {"page_number": number})
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table

        document = Document(io.BytesIO(data))
        headings = []
        table_number = 0
        for index, element in enumerate(document.iter_inner_content(), 1):
            if isinstance(element, Table):
                table_number += 1
                for row_index, row in enumerate(element.rows, 1):
                    for column, cell in enumerate(row.cells, 1):
                        add(
                            cell.text,
                            {
                                "section_path": list(headings),
                                "table_number": table_number,
                                "row": row_index,
                                "column": column,
                            },
                        )
            else:
                style = element.style.name if element.style else ""
                if style.startswith("Heading ") and style[8:].isdigit():
                    level = int(style[8:])
                    headings = headings[: max(level - 1, 0)] + [element.text]
                add(element.text, {"paragraph_number": index, "section_path": list(headings)})
    elif extension in {".html", ".htm"}:
        from app.webfetch import decode_html

        walker = _HtmlBlocks(add)
        walker.feed(decode_html(data))
        walker.close()
        title = " ".join(walker.title.split())[:300] or None
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) > 50:
                raise ParseIssue("sheet_limit")
            for sheet in workbook:
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                    raise ParseIssue("spreadsheet_dimensions_limit")
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            add(
                                cell.value,
                                {
                                    "sheet_name": sheet.title,
                                    "cell_range": cell.coordinate,
                                    "formula_not_evaluated": cell.data_type == "f",
                                },
                            )
        finally:
            workbook.close()
    return {
        "blocks": blocks,
        "status": "needs_attention" if issue or not blocks else "parsed",
        "error_code": issue or ("no_extractable_text" if not blocks else None),
        "title": title,
    }


def _child(connection, filename, data, transcript):
    # Do not emit parser diagnostics containing source content to routine logs.
    logging.disable(logging.CRITICAL)
    try:
        try:
            import resource

            resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
        except ImportError:  # Windows local development; production worker uses Linux.
            pass
        connection.send(parse_document(filename, data, transcript))
    except ParseIssue as error:
        connection.send({"blocks": [], "status": "needs_attention", "error_code": str(error)})
    except Exception:
        connection.send({"blocks": [], "status": "failed", "error_code": "parse_failed"})
    finally:
        connection.close()


def parse_isolated(filename: str, data: bytes, transcript=False) -> dict:
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(target=_child, args=(child, filename, data, transcript))
    process.start()
    child.close()
    try:
        if parent.poll(settings.parse_timeout_seconds):
            try:
                return parent.recv()
            except EOFError:
                return {"blocks": [], "status": "failed", "error_code": "parser_resource_limit"}
        return {"blocks": [], "status": "failed", "error_code": "parse_timeout"}
    finally:
        if process.is_alive():
            process.terminate()
        process.join(timeout=5)
        parent.close()
