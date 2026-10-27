"""M1 behavior at the HTTP and storage boundaries, using disposable databases."""

import io
from datetime import datetime, timedelta, timezone

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from sqlalchemy import func, select

from app import worker
from app.models import AuditEvent, Job, ProjectMembership, SourceBlock, SourceVersion
from app.parsing import parse_document, parse_isolated
from app.seed import key
from tests.support import ORIGIN, post, sign_in


def test_login_logout_csrf_and_rate_limit(client):
    browser, factory, _ = client
    assert browser.get("/api/projects").status_code == 401
    assert (
        post(
            browser, "/api/auth/login", {"email": "specialist@example.invalid", "password": "wrong"}
        ).status_code
        == 401
    )
    csrf = sign_in(browser)
    assert browser.get("/api/auth/me").json()["role"] == "admin"
    assert (
        browser.post("/api/organizations", json={"name": "Rejected"}, headers=ORIGIN).status_code
        == 403
    )
    assert (
        browser.post(
            "/api/organizations", json={"name": "Rejected"}, headers={"x-csrf-token": csrf}
        ).status_code
        == 403
    )
    result = post(browser, "/api/organizations", {"name": "Created"}, csrf)
    assert result.status_code == 201
    assert post(browser, "/api/auth/logout", {}, csrf).status_code == 204
    assert browser.get("/api/projects").status_code == 401
    with factory() as db:
        assert (
            db.scalar(select(func.count()).select_from(AuditEvent)) == 0
        )  # Organization is not project bound.


def test_project_membership_isolation_across_routes(client):
    browser, factory, memory = client
    csrf = sign_in(browser)
    org = post(browser, "/api/organizations", {"name": "New group"}, csrf).json()
    project = post(
        browser, "/api/projects", {"name": "New project", "organization_id": org["id"]}, csrf
    ).json()
    pid = project["id"]
    portfolio = post(browser, f"/api/projects/{pid}/portfolios", {"name": "East"}, csrf).json()
    prop = post(
        browser,
        f"/api/projects/{pid}/properties",
        {"name": "Building", "portfolio_id": portfolio["id"]},
        csrf,
    )
    assert prop.status_code == 201
    contact = post(
        browser,
        f"/api/projects/{pid}/contacts",
        {"name": "Customer", "email": "c@example.invalid", "role": "Manager"},
        csrf,
    )
    assert contact.status_code == 201
    upload = browser.post(
        f"/api/projects/{pid}/sources/upload",
        files={"file": ("rules.txt", b"First rule\nSecond rule", "text/plain")},
        headers={**ORIGIN, "x-csrf-token": csrf},
    )
    assert upload.status_code == 201, upload.text
    version = upload.json()
    assert version["content_hash"] and version["has_original"]
    assert len(memory.objects) == 1
    with factory() as db:
        assert db.scalar(
            select(ProjectMembership).where(
                ProjectMembership.project_id == pid, ProjectMembership.user_id == key("staff")
            )
        )
    post(browser, "/api/auth/logout", {}, csrf)
    other_csrf = sign_in(browser, "outsider@example.invalid", "another-fictional-passphrase")
    assert pid not in {p["id"] for p in browser.get("/api/projects").json()}
    for path in (
        f"/api/projects/{pid}",
        f"/api/projects/{pid}/scopes",
        f"/api/projects/{pid}/contacts",
        f"/api/projects/{pid}/sources",
        f"/api/projects/{pid}/sources/versions/{version['id']}/blocks",
        f"/api/projects/{pid}/sources/versions/{version['id']}/original",
    ):
        assert browser.get(path).status_code == 404, path
    assert (
        post(
            browser, f"/api/projects/{pid}/properties", {"name": "Intrusion"}, other_csrf
        ).status_code
        == 404
    )
    assert (
        browser.post(
            f"/api/projects/{pid}/sources/upload",
            files={"file": ("rules.txt", b"Intrusion")},
            headers={**ORIGIN, "x-csrf-token": other_csrf},
        ).status_code
        == 404
    )
    assert not browser.get("/api/organizations").json()


def test_private_upload_worker_evidence_and_versions(client, monkeypatch):
    browser, factory, memory = client
    monkeypatch.setattr(worker, "SessionLocal", factory)
    csrf = sign_in(browser)
    pid = key("project")
    uploaded = browser.post(
        f"/api/projects/{pid}/sources/upload",
        files={"file": ("sop.txt", b"Pets: two cats\nWeekend: call supervisor", "text/plain")},
        headers={**ORIGIN, "x-csrf-token": csrf},
    )
    assert uploaded.status_code == 201, uploaded.text
    v1 = uploaded.json()
    assert v1["parse_status"] == "pending"
    assert worker.run_once(memory, parse_document)
    evidence = browser.get(f"/api/projects/{pid}/sources/versions/{v1['id']}/blocks").json()
    assert evidence["version"]["parse_status"] == "parsed"
    assert evidence["total"] == 2
    assert evidence["blocks"][0]["locator"] == {"line_start": 1, "line_end": 1}
    original = browser.get(f"/api/projects/{pid}/sources/versions/{v1['id']}/original")
    assert original.content == b"Pets: two cats\nWeekend: call supervisor"
    assert original.headers["content-type"] == "application/octet-stream"
    v2 = browser.post(
        f"/api/projects/{pid}/sources/upload?document_id={v1['document_id']}",
        files={"file": ("sop.txt", b"Pets: one cat", "text/plain")},
        headers={**ORIGIN, "x-csrf-token": csrf},
    ).json()
    assert v2["version"] == 2 and v2["content_hash"] != v1["content_hash"]
    assert worker.run_once(memory, parse_document)
    versions = browser.get(f"/api/projects/{pid}/sources").json()
    target = next(source for source in versions if source["id"] == v1["document_id"])
    assert [v["version"] for v in target["versions"]] == [2, 1]
    assert (
        browser.get(f"/api/projects/{pid}/sources/versions/{v1['id']}/original").content
        == original.content
    )
    with factory() as db:
        assert (
            db.scalar(
                select(func.count())
                .select_from(SourceBlock)
                .where(SourceBlock.source_version_id == v1["id"])
            )
            == 2
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(Job)
                .where(Job.project_id == pid, Job.kind == "ingestion")
            )
            == 2
        )


def test_transcript_timecodes_and_upload_validation(client, monkeypatch):
    browser, factory, memory = client
    monkeypatch.setattr(worker, "SessionLocal", factory)
    csrf = sign_in(browser)
    pid = key("project")
    transcript = post(
        browser,
        f"/api/projects/{pid}/sources/transcript",
        {"name": "Call", "text": "00:12 Intro\n01:03 Emergency transfer"},
        csrf,
    )
    assert transcript.status_code == 201
    assert worker.run_once(memory, parse_document)
    blocks = browser.get(
        f"/api/projects/{pid}/sources/versions/{transcript.json()['id']}/blocks"
    ).json()["blocks"]
    assert [block["locator"]["timestamp_start"] for block in blocks] == [12, 63]
    for name, data in (("bad.pdf", b"not a pdf"), ("bad.exe", b"bad"), ("empty.txt", b"")):
        response = browser.post(
            f"/api/projects/{pid}/sources/upload",
            files={"file": (name, data)},
            headers={**ORIGIN, "x-csrf-token": csrf},
        )
        assert response.status_code == 422
    from app.config import settings

    monkeypatch.setattr(settings, "max_upload_bytes", 16)
    too_big = browser.post(
        f"/api/projects/{pid}/sources/upload",
        files={"file": ("large.txt", b"A" * 17)},
        headers={**ORIGIN, "x-csrf-token": csrf},
    )
    assert too_big.status_code == 413


def test_pdf_docx_xlsx_locators_and_scanned_attention():
    pdf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf)
    result = parse_document("scan.pdf", pdf.getvalue())
    assert result["status"] == "needs_attention" and result["blocks"] == []
    doc = Document()
    doc.add_heading("Pets", 1)
    doc.add_paragraph("Two cats permitted")
    docx = io.BytesIO()
    doc.save(docx)
    result = parse_document("guide.docx", docx.getvalue())
    assert any("Pets" in block["locator"]["section_path"] for block in result["blocks"])
    assert all("page_number" not in block["locator"] for block in result["blocks"])
    workbook = Workbook()
    workbook.active.title = "Rules"
    workbook.active["B3"] = "Weekend route"
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    result = parse_document("roster.xlsx", xlsx.getvalue())
    assert result["blocks"] == [
        {
            "text": "Weekend route",
            "locator": {"sheet_name": "Rules", "cell_range": "B3", "formula_not_evaluated": False},
            "position": 0,
        }
    ]


def test_parser_process_is_bounded():
    result = parse_isolated("rules.txt", b"A rule")
    assert result["status"] == "parsed"


def test_abandoned_job_recovered(client, monkeypatch):
    browser, factory, memory = client
    monkeypatch.setattr(worker, "SessionLocal", factory)
    csrf = sign_in(browser)
    pid = key("project")
    v = browser.post(
        f"/api/projects/{pid}/sources/upload",
        files={"file": ("rules.txt", b"Recovered", "text/plain")},
        headers={**ORIGIN, "x-csrf-token": csrf},
    ).json()
    with factory.begin() as db:
        job = db.scalar(select(Job).where(Job.source_version_id == v["id"]))
        job.status = "running"
        job.claimed_at = datetime.now(timezone.utc) - timedelta(minutes=3)
        job.attempts = 1
    assert worker.run_once(memory, parse_document)
    with factory() as db:
        job = db.scalar(select(Job).where(Job.source_version_id == v["id"]))
        assert job.status == "parsed" and job.attempts == 2
        assert db.get(SourceVersion, v["id"]).parse_status == "parsed"
